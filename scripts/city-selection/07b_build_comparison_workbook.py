# Step 7b · city-selection · build news_coverage_ala_comparison.xlsx.
# Script as run on 2026-09-27 (originally build.py). Reads Gina's Numbers file from its iCloud path and city_results.jsonl (Step 7a).
# Usage: python3 07b_build_comparison_workbook.py <output.xlsx>
import json, sys, csv, urllib.parse, statistics
from numbers_parser import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OUT = sys.argv[1]
src = Document("/Users/ginah/Library/Mobile Documents/com~apple~Numbers/Documents/news_coverage_by_state_cross_pm.numbers").sheets[0].tables[0]
rows = list(src.rows(values_only=True)); hdr = rows[0]; srows = [dict(zip(hdr, r)) for r in rows[1:]]
cities = [json.loads(l) for l in open('city_results.jsonl')]
flag = {'Springfield':'Also a city in IL/MA/MO','College':'Common word','Pasadena':'Also a city in CA','Warren':'Also a surname (e.g., Elizabeth Warren)','Mesa':'Also a common word / other places','Carmel':'Also Carmel, CA','Delano':'Also a surname'}
CNAME = {r['collection_id_cs']: r['collection_name'] for r in srows}
CSRC = {r['collection_id_cs']: r['total_sources_in_collection'] for r in srows}

def srow(metro, st):
    m = [r for r in srows if r['metro_area'] == metro and r['state_abbr'] == st]
    assert len(m) == 1, (metro, st); return m[0]

ARIAL = 'Arial'
f = lambda **k: Font(name=ARIAL, **k)
thin = Side(style='thin', color='BFBFBF')
FILL = {'id': 'D9E1F2', 'ala': 'FCE4D6', 'state': 'E2EFDA', 'city': 'FFF2CC', 'cmp': 'EDE1F5', 'ref': 'EDEDED'}

wb = Workbook()
# ---------- Method sheet (thresholds) ----------
me = wb.active; me.title = 'Method'
# ---------- States sheet ----------
ss = wb.create_sheet('States')
sh = ['state_abbr','State','query_mediacloud','collection_id_cs','collection_name','total_sources_in_collection','total_stories','index_stories_sources']
ss.append(sh)
seen = []
for r in srows:
    if r['state_abbr'] in seen: continue
    seen.append(r['state_abbr'])
    ss.append([r['state_abbr'], r['State'], r['query_mediacloud'], int(r['collection_id_cs']), r['collection_name'], int(r['total_sources_in_collection']), int(r['total_stories']), None])
NS = len(seen) + 1
for i in range(2, NS + 1):
    ss[f'H{i}'] = f'=G{i}/F{i}'; ss[f'H{i}'].number_format = '0.00'
# ---------- Comparison sheet ----------
ws = wb.create_sheet('Comparison', 0)
cols = [
 ('#','id'),('city','id'),('state','id'),('metro_area','id'),
 ('ala_shortterm_pm25_rank_2026','ala'),('ala_yearlong_pm25_rank_2026','ala'),('pollution_score (avg of ALA ranks)','ala'),('pollution_group','ala'),
 ('state_query','state'),('state_collection_id_cs','state'),('state_collection_name','state'),('state_total_sources','state'),('state_total_stories','state'),('state_index (stories/sources)','state'),('state_coverage_rank','state'),('state_coverage_group','state'),
 ('city_query','city'),('city_collection_id_cs','city'),('city_collection_name','city'),('city_total_sources','city'),('city_total_stories','city'),('city_index (stories/sources)','city'),('city_coverage_rank','city'),('city_coverage_group','city'),
 ('rank_change (state rank − city rank)','cmp'),('quadrant_state_level','cmp'),('quadrant_city_level','cmp'),('quadrant_changed?','cmp'),
 ('city_ambiguity_note','ref'),('state_API_request_URL','ref'),('city_API_request_URL','ref')]
ws.append([c for c, _ in cols])
N = len(cities) + 1
for i, c in enumerate(cities, 2):
    s = srow(c['metro'], c['state'])
    url = 'https://search.mediacloud.org/api/search/total-count?' + urllib.parse.urlencode(dict(q=c['q'], start='2024-09-25', end='2026-09-25', cs=c['cs'], platform='onlinenews-mediacloud'))
    st_short = int(s['ala_shortterm_pm25_rank_2026']) if s['ala_shortterm_pm25_rank_2026'] is not None else None
    ws.append([i-1, c['city'], c['state'], c['metro'], st_short, int(s['ala_yearlong_pm25_rank_2026']),
      f'=AVERAGE(E{i}:F{i})', f'=IF(G{i}<=Method!$B$4,"High pollution","Low pollution")',
      s['query_mediacloud'], int(s['collection_id_cs']), s['collection_name'], int(s['total_sources_in_collection']), int(s['total_stories']),
      f'=M{i}/L{i}', f'=RANK(N{i},$N$2:$N${N},0)', f'=IF(N{i}>=Method!$B$5,"High coverage","Low coverage")',
      c['q'], c['cs'], CNAME[c['cs']], int(CSRC[c['cs']]), c['result']['count']['relevant'],
      f'=U{i}/T{i}', f'=RANK(V{i},$V$2:$V${N},0)', f'=IF(V{i}>=Method!$B$6,"High coverage","Low coverage")',
      f'=O{i}-W{i}', f'=H{i}&" / "&P{i}', f'=H{i}&" / "&X{i}', f'=IF(Z{i}=AA{i},"No","Yes")',
      flag.get(c['city'], ''), s['API_request_URL_claude'], url])

for j, (name, grp) in enumerate(cols, 1):
    cell = ws.cell(1, j); cell.font = f(bold=True); cell.fill = PatternFill('solid', fgColor=FILL[grp])
    cell.alignment = Alignment(wrap_text=True, vertical='center'); cell.border = Border(bottom=thin)
    for i in range(2, N + 1):
        x = ws.cell(i, j); x.font = f(bold=(grp in ('id', 'ala') and j in (2, 3, 5, 6)))
        if grp in ('id', 'ala') and j in (2, 3, 5, 6): x.fill = PatternFill('solid', fgColor=FILL[grp])
    ws.column_dimensions[get_column_letter(j)].width = {'#':5,'city':15,'state':7,'metro_area':34}.get(name, 48 if 'URL' in name else 40 if 'query' in name else 36 if 'collection_name' in name else 26 if 'quadrant' in name else 16)
for i in range(2, N + 1):
    for col in ('G', 'N', 'V'): ws[f'{col}{i}'].number_format = '0.00'
    ws[f'E{i}'].comment = None
ws.row_dimensions[1].height = 45
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{N}'
# highlight quadrants and changes
QC = {'High pollution / Low coverage': 'F8CBAD', 'High pollution / High coverage': 'C6EFCE', 'Low pollution / High coverage': 'BDD7EE', 'Low pollution / Low coverage': 'EDEDED'}
for rng in (f'Z2:Z{N}', f'AA2:AA{N}'):
    c0 = rng.split(':')[0][:-1]
    for k, v in QC.items():
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{c0}2="{k}"'], fill=PatternFill('solid', fgColor=v)))
ws.conditional_formatting.add(f'AB2:AB{N}', FormulaRule(formula=['AB2="Yes"'], fill=PatternFill('solid', fgColor='FFC7CE'), font=Font(name=ARIAL, bold=True, color='9C0006')))
ws['E1'].comment = Comment('Source: American Lung Association, State of the Air 2026, as provided by the user (news_coverage_by_state_cross_pm.numbers). Blank = not ranked in short-term list; pollution_score then uses year-round rank only.', 'Claude')

# ---------- Method content ----------
me['A1'] = 'Method & thresholds'; me['A1'].font = f(bold=True, size=13)
me['A3'] = 'Threshold'; me['B3'] = 'Value'; me['C3'] = 'Definition'
for c in ('A3', 'B3', 'C3'): me[c].font = f(bold=True)
me['A4'] = 'Pollution threshold'; me['B4'] = f'=MEDIAN(Comparison!G2:G{N})'
me['C4'] = 'Median pollution_score across all 36 city rows. pollution_score = average of ALA short-term and year-round PM2.5 ranks (lower = worse air). High pollution = score ≤ threshold.'
me['A5'] = 'State coverage threshold'; me['B5'] = f'=MEDIAN(States!H2:H{NS})'
me['C5'] = 'Median state_index across the 10 state collections. High coverage = state_index ≥ threshold.'
me['A6'] = 'City coverage threshold'; me['B6'] = f'=MEDIAN(Comparison!V2:V{N})'
me['C6'] = 'Median city_index across the 36 city queries. High coverage = city_index ≥ threshold.'
for c in ('B4', 'B5', 'B6'): me[c].number_format = '0.00'; me[c].font = f()
notes = ['Notes',
 'State-level data: copied from the user-supplied file news_coverage_by_state_cross_pm.numbers (query "air pollution" OR "air quality" in each state collection).',
 'City-level data: Media Cloud API, query ("air pollution" OR "air quality") AND "[city]" in the city\'s own state collection, 2024-09-25 to 2026-09-25, platform onlinenews-mediacloud.',
 'Pittsburgh-Weirton-Steubenville: Pittsburgh searched in Pennsylvania, Weirton in West Virginia, Steubenville in Ohio; each row uses that state\'s state-level data.',
 'San Diego-Chula Vista-Carlsbad and Houston-Pasadena have no ALA short-term rank; their pollution_score uses the year-round rank only.',
 'Ranks: 1 = highest index. Ties share a rank (state-level ranks tie for every city in the same state). rank_change > 0 means the city ranks higher on city-level mentions than on state-level mentions.',
 'Ambiguous city names (see city_ambiguity_note) may inflate city-level counts.',
 'Column colors: blue = identity, orange = ALA rankings, green = state-level coverage, yellow = city-level coverage, purple = comparison, grey = reference.']
for k, t in enumerate(notes, 8):
    me[f'A{k}'] = t; me[f'A{k}'].font = f(bold=(k == 8))
for c in ('A4', 'A5', 'A6', 'C4', 'C5', 'C6'): me[c].font = f()
me.column_dimensions['A'].width = 26; me.column_dimensions['B'].width = 10; me.column_dimensions['C'].width = 120
for j in range(1, 9):
    ss.cell(1, j).font = f(bold=True); ss.cell(1, j).fill = PatternFill('solid', fgColor=FILL['state'])
    ss.column_dimensions[get_column_letter(j)].width = 40 if j in (3, 5) else 16
    for i in range(2, NS + 1): ss.cell(i, j).font = f()
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT)
